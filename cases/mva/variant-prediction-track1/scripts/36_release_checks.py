"""Check Track 1 release constraints omitted by validate_submission.py.

Usage: python scripts/36_release_checks.py [csv] [report] [workbook]
No patient VCF or genome-wide file is loaded.
"""

from __future__ import annotations

import csv
import re
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
TEAM = "Nexgene-AI"
CSV = ROOT / f"submission/{TEAM}_mn5-evidence-v2.csv"
REPORT = ROOT / f"submission/{TEAM}_track1_report.md"
BOOK = ROOT / f"submission/{TEAM}_methods_description.xlsx"
PROBAND = "PROBAND01"  # Official scorer submit_track1.py, verified 2026-10-16.


def main(argv: list[str]) -> int:
    paths = [Path(arg) for arg in argv] if argv else [CSV, REPORT, BOOK]
    if len(paths) != 3:
        raise ValueError("expected CSV, report and methods workbook paths")
    csv_file, report, book = paths
    issues: list[str] = []
    for path in paths:
        if not path.is_file():
            issues.append(f"missing {path}")
    if not re.fullmatch(rf"{re.escape(TEAM)}_[A-Za-z0-9-]+\.csv", csv_file.name):
        issues.append("CSV must follow <team>_<approach>.csv")
    if report.name != f"{TEAM}_track1_report.md":
        issues.append("report must be <team>_track1_report.md")
    if book.name != f"{TEAM}_methods_description.xlsx":
        issues.append("methods workbook name does not match expected name")
    if csv_file.is_file():
        with csv_file.open(newline="", encoding="utf-8-sig") as stream:
            rows = list(csv.DictReader(stream))
        if not rows or any(row["proband_id"] != PROBAND for row in rows):
            issues.append("proband_id must match official scorer PROBAND01")
        if any(row["finding_type"] == "secondary" for row in rows):
            issues.append("review any incidental findings manually")
        for row in rows:
            if "phase=" not in row["notes"] or "limitation=" not in row["notes"]:
                issues.append("notes must record phase and principal limitation")
                break
    if report.is_file():
        text = report.read_text(encoding="utf-8").lower()
        for label in (
            "ai disclosure",
            "arf machine",
            "asa-w1",
            "funding",
            "ehpc-aif-2026pg01-1286",
            "marenostrum 5",
            "not a clinical diagnostic report",
            "phase",
        ):
            if label not in text:
                issues.append(f"report missing {label!r}")
    if book.is_file():
        from openpyxl import load_workbook

        wb = load_workbook(book, read_only=True)
        ws = wb["Track 1 methods"]
        for line in range(7, 20):
            if ws[f"A{line}"].value and not str(ws[f"B{line}"].value or "").strip():
                issues.append(f"methods form B{line} unanswered")
    print(
        "\n".join(issues)
        if issues
        else f"OK: release checks; {len(rows)} rows; {PROBAND}"
    )
    return int(bool(issues))


if __name__ == "__main__":
    try:
        sys.exit(main(sys.argv[1:]))
    except (OSError, KeyError, ValueError) as exc:
        sys.exit(f"ERROR: {exc}")
