"""Fill a layout-preserving copy of the official Track 1 methods workbook."""

from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook

import importlib.util

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "submission/methods_description_form.xlsx"
TARGET = ROOT / "submission/Nexgene-AI_methods_description.xlsx"
REPORT = ROOT / "submission/Nexgene-AI_track1_report.md"


def main() -> None:
    if TARGET.exists():
        raise FileExistsError(TARGET)
    if not REPORT.is_file():
        raise FileNotFoundError(REPORT)
    wb = load_workbook(SOURCE)
    ws = wb["Track 1 methods"]
    answers = {
        7: "Nexgene AI",
        11: "Hybrid: initial computational rankings automated; strict manual adjudication of molecular mechanism, phase, eligibility and source independence before final CSV.",
        12: "Read-backing, mapping/paralog and transcript interpretation, source-to-allele matching, phenotype specificity, AF exceptions, causality and red-team objections; equal heuristic scores resolved by frozen multi-factor tie hierarchy without altering ePCR.",
        14: (
            "GRCh38 no-alt reference, gnomAD allele frequencies, ClinVar assertions/review status, MANE/snpEff consequences, "
            "HPO/gene-disease sources, independent published case series; query identifiers and dates are in "
            "results/evidence/literature/literature_queries.tsv and the report reference list."
        ),
        16: "Both: compound-heterozygote pairs are represented by two alleles on one CSV row; single-variant hypotheses are permitted when inheritance warrants. Two 0/1 genotypes are not called trans without physical/familial evidence.",
        17: "Only genuinely unrelated incidental findings qualify as secondary. Competing causal hypotheses are primary; no incidental result is included without separate support.",
        18: (
            "Alignment, three-caller genotyping, annotation, phasing, CNV/SV, SpliceAI and mosaic screens ran as MareNostrum 5 "
            "Slurm jobs (Appendix A3 of submission/Nexgene-AI_track1_report.md; job identifiers also in docs/lineage.md). "
            "Ranking, ACMG/AMP adjudication and the report were completed on a local workstation. No commercial generative-AI "
            "API was used (ASA-W1 ran locally), so that spend is zero. MN5 allocation was the EuroHPC award "
            "EHPC-AIF-2026PG01-1286; a single wall-clock total was not recorded."
        ),
    }
    spec = importlib.util.spec_from_file_location("methods_answers", ROOT / "scripts/55_update_methods_form.py")
    current = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(current)
    answers.update(current.ANSWERS)
    for row, answer in answers.items():
        assert ws[f"A{row}"].value
        ws[f"B{row}"] = answer
    wb.save(TARGET)
    print(f"Filled {len(answers)} Track 1 fields: {TARGET}")


if __name__ == "__main__":
    main()
